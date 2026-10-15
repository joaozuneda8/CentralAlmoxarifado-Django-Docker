"""Leitura do .xlsx exportado das planilhas do Apps Script (sem ORM, sem regras de negócio).

`ler_xlsx` devolve as linhas de cada aba como dicionários cabeçalho → valor da célula; a
conversão dos valores (`texto_celula`, `data_celula`, `momento_celula`, `qtd_celula`) fica
para quem usa a leitura, que sabe o que cada coluna deveria conter.
"""

import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from django.utils import timezone

from pedidos import regras

Linha = dict[str, object]

# Colunas obrigatórias por aba, com os nomes exatos de `Config.gs` (constantes COL_).
ABAS: dict[str, tuple[str, ...]] = {
    "Pedidos": (
        "Nº",
        "ID",
        "Origem",
        "Criado em",
        "Solicitante",
        "Destino",
        "Observação",
        "Atendimento_ID",
        "Data efetiva",
        "Recebedor",
        "Centro de custo",
        "Tipo Saída",
        "Placa",
        "Obs. da saída",
        "Atendido em",
    ),
    "Itens": (
        "Nº",
        "ID",
        "Código",
        "Material",
        "Unidade",
        "Categoria",
        "Grupo",
        "Subgrupo",
        "Estoque no pedido",
        "Qtd pedida",
        "Qtd entregue",
    ),
    "Saídas": (
        "Data",
        "Qtd",
        "Código",
        "Descrição do Material",
        "Recebimento",
        "Centro de Custo",
        "Tipo Saída",
        "Veículo (Placa)",
        "Observações",
        "Requisição",
        "Status SCPI",
        "Lançado SCPI em",
        "Origem_ID",
        "Lote_SCPI",
    ),
    "Centros de Custo": ("DESCR", "CODCCUSTO", "PLACA"),
    "Listas": ("Tipos de saída", "Grupos ocultos (AB)", "Categorias (limpeza)", "Funcionários"),
    "Lista_Limpeza": ("Código", "Categoria", "Nome exibido", "Unidade exibida"),
}
# Colunas que podem faltar: "Recibo impresso em" (Pedidos) e "Destinos" (Listas) são mais novas.
# Quando existem, vêm na linha como qualquer outra.

_DATA_BR = re.compile(r"^(\d{2})/(\d{2})/(\d{4})$")
_DATA_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")
_MOMENTO_BR = re.compile(r"^(\d{2})/(\d{2})/(\d{4})(?: (\d{2}):(\d{2})(?::(\d{2}))?)?$")


class ErroPlanilha(ValueError):
    """A planilha não tem a forma esperada (aba ou coluna faltando)."""


def texto_celula(valor: object) -> str:
    """Texto da célula: `None` → "", inteiro sem casas, sem espaços nas pontas nem `'` inicial."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, str):
        texto = valor.strip()
        # `sanitizarTexto_` do Apps Script põe um apóstrofo na frente do que parece fórmula.
        return texto[1:] if texto.startswith("'") else texto
    return str(valor)


def data_celula(valor: object) -> date | None:
    """Data da célula: `datetime`/`date` ou texto `dd/mm/aaaa` ou `aaaa-mm-dd`; vazio → `None`."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, str):
        texto = valor.strip()
        if m := _DATA_BR.match(texto):
            dia, mes, ano = m.groups()
        elif m := _DATA_ISO.match(texto):
            ano, mes, dia = m.groups()
        else:
            raise ValueError(f"Data inválida: {valor!r}.")
        try:
            return date(int(ano), int(mes), int(dia))
        except ValueError:
            raise ValueError(f"Data inválida: {valor!r}.") from None
    raise ValueError(f"Data inválida: {valor!r}.")


def momento_celula(valor: object) -> datetime | None:
    """Momento da célula, no fuso `TIME_ZONE`: `datetime` ingênuo ou `dd/mm/aaaa[ hh:mm[:ss]]`."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, datetime):
        momento = valor
    elif isinstance(valor, date):
        momento = datetime(valor.year, valor.month, valor.day)
    elif isinstance(valor, str) and (m := _MOMENTO_BR.match(valor.strip())):
        dia, mes, ano, hora, minuto, segundo = m.groups()
        try:
            momento = datetime(
                int(ano), int(mes), int(dia), int(hora or 0), int(minuto or 0), int(segundo or 0)
            )
        except ValueError:
            raise ValueError(f"Momento inválido: {valor!r}.") from None
    else:
        raise ValueError(f"Momento inválido: {valor!r}.")
    return timezone.make_aware(momento) if timezone.is_naive(momento) else momento


def qtd_celula(valor: object) -> Decimal | None:
    """Quantidade da célula: número ou texto pt-BR ("1,5"); vazio → `None`; negativo → erro."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, bool):
        raise ValueError(f"Quantidade inválida: {valor!r}.")
    if isinstance(valor, int | float):
        try:
            qtd = Decimal(repr(valor)).quantize(Decimal("0.001"))
        except InvalidOperation:
            raise ValueError(f"Quantidade inválida: {valor!r}.") from None
    else:
        qtd = regras.parse_qtd_br(valor)
        if qtd is None:
            raise ValueError(f"Quantidade inválida: {valor!r}.")
    if qtd < 0:
        raise ValueError(f"Quantidade inválida: {valor!r}.")
    return qtd


def _vazia(valor: object) -> bool:
    return valor is None or (isinstance(valor, str) and valor == "")


def ler_xlsx(
    caminho: Path, abas: dict[str, tuple[str, ...]] | None = None
) -> dict[str, list[Linha]]:
    """Lê as abas pedidas (padrão: `ABAS`), e só elas, com as colunas obrigatórias de cada uma.
    Cada linha leva o número real na chave `_linha`."""
    pedidas = ABAS if abas is None else abas
    livro = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        problemas = [
            f'A planilha não tem a aba "{nome}".'
            for nome in pedidas
            if nome not in livro.sheetnames
        ]
        abas: dict[str, list[Linha]] = {}
        for nome, obrigatorias in pedidas.items():
            if nome not in livro.sheetnames:
                continue
            linhas = livro[nome].iter_rows(values_only=True)
            cabecalho = next(linhas, ())
            colunas: dict[str, int] = {}
            for posicao, celula in enumerate(cabecalho):
                titulo = texto_celula(celula)
                if titulo and titulo not in colunas:
                    colunas[titulo] = posicao
            problemas += [
                f'{nome}: coluna "{coluna}" não encontrada.'
                for coluna in obrigatorias
                if coluna not in colunas
            ]
            abas[nome] = [
                {
                    **{
                        titulo: (celulas[posicao] if posicao < len(celulas) else None)
                        for titulo, posicao in colunas.items()
                    },
                    "_linha": numero,
                }
                for numero, celulas in enumerate(linhas, start=2)
                if not all(_vazia(celula) for celula in celulas)
            ]
    finally:
        livro.close()
    if problemas:
        raise ErroPlanilha("\n".join(problemas))
    return abas
